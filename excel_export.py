import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime

# <<< THÊM MỚI: HÀM XUẤT EXCEL BÁO CÁO THÁNG >>>
def generate_monthly_summary_excel(student_data_list, month, year, teacher_name, class_name, filepath):
    """
    Tạo file Excel tổng hợp kết quả rèn luyện theo tháng, bao gồm sheet xếp hạng tổ.
    """
    try:
        wb = openpyxl.Workbook()
        
        # --- Định nghĩa các style chung ---
        font_title = Font(name='Times New Roman', size=16, bold=True)
        font_header = Font(name='Times New Roman', size=12, bold=True)
        font_bold = Font(name='Times New Roman', size=11, bold=True)
        font_normal = Font(name='Times New Roman', size=11)
        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        header_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        # === BẮT ĐẦU SHEET 1: TỔNG KẾT THEO HỌC SINH ===
        ws = wb.active
        ws.title = f"Tong_Ket_Thang_{month}_{year}"

        # --- Header của file ---
        ws.merge_cells('A1:E1'); ws['A1'].value = f"BÁO CÁO TỔNG KẾT RÈN LUYỆN THÁNG {month}/{year}"; ws['A1'].font = font_title; ws['A1'].alignment = center_alignment
        ws.merge_cells('A2:E2'); ws['A2'].value = f"LỚP: {class_name}"; ws['A2'].font = font_header; ws['A2'].alignment = center_alignment
        ws.merge_cells('A3:E3'); ws['A3'].value = f"GVCN: {teacher_name}"; ws['A3'].font = font_header; ws['A3'].alignment = center_alignment
        ws.append([])

        # --- Header của bảng dữ liệu ---
        headers = ['STT', 'Họ và Tên', 'Tổ', 'Điểm Tháng', 'Xếp Loại Tháng']
        ws.append(headers)
        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.font = font_bold; cell.alignment = center_alignment; cell.fill = header_fill; cell.border = thin_border

        # --- Đổ dữ liệu học sinh ---
        for i, student_data in enumerate(student_data_list, 1):
            row_data = [i, student_data['ten'], student_data.get('to') or '', student_data['diem'], student_data['xeploai']]
            ws.append(row_data)
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=ws.max_row, column=col_idx); cell.font = font_normal; cell.border = thin_border
                if col_idx in [1, 3, 4, 5]: cell.alignment = center_alignment
                else: cell.alignment = left_alignment
        
        # --- Tùy chỉnh độ rộng cột và in ấn ---
        ws.column_dimensions['A'].width = 5; ws.column_dimensions['B'].width = 35; ws.column_dimensions['C'].width = 10
        ws.column_dimensions['D'].width = 15; ws.column_dimensions['E'].width = 20
        ws.print_options.horizontalCentered = True; ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        
        # === BẮT ĐẦU SHEET 2: TỔNG KẾT TỔ THÁNG (THÊM MỚI) ===
        ws_rank = wb.create_sheet("Xep_Hang_To_Thang")
        teams = {}; team_ranking = []
        if student_data_list:
            for student in student_data_list:
                team_name = student.get('to')
                if team_name and team_name.strip():
                    if team_name not in teams: teams[team_name] = {'total_score': 0, 'count': 0}
                    teams[team_name]['total_score'] += student['diem']; teams[team_name]['count'] += 1
        for name, data in teams.items():
            if data['count'] > 0: avg_score = round(data['total_score'] / data['count'], 2); team_ranking.append({'name': name, 'avg_score': avg_score})
        
        team_ranking.sort(key=lambda x: x['avg_score'], reverse=True)

        ws_rank.merge_cells('A1:C1'); ws_rank['A1'].value = f"BẢNG XẾP HẠNG THI ĐUA TỔ - THÁNG {month}/{year}"; ws_rank['A1'].font = font_title; ws_rank['A1'].alignment = center_alignment; ws_rank.row_dimensions[1].height = 30
        ws_rank.append([])
        rank_headers = ['Xếp Hạng', 'Tên Tổ', 'Điểm Trung Bình']; ws_rank.append(rank_headers)
        for cell in ws_rank[3]: cell.font = font_bold; cell.alignment = center_alignment; cell.fill = header_fill; cell.border = thin_border
        
        if not team_ranking:
             ws_rank.append(['', 'Không có dữ liệu của tổ để xếp hạng.'])
        else:
            for idx, team in enumerate(team_ranking, 1):
                ws_rank.append([idx, f"Tổ {team['name']}", team['avg_score']])
                for col_idx in range(1, 4):
                    cell = ws_rank.cell(row=ws_rank.max_row, column=col_idx); cell.font = font_normal; cell.alignment = center_alignment; cell.border = thin_border
        
        ws_rank.column_dimensions['A'].width = 15; ws_rank.column_dimensions['B'].width = 25; ws_rank.column_dimensions['C'].width = 25
        ws_rank.print_options.horizontalCentered = True
        
        wb.active = 0 # Mở file Excel sẽ hiện sheet đầu tiên
        wb.save(filepath)
        return True, "Xuất file Excel thành công."
    except Exception as e:
        return False, f"Lỗi khi tạo file Excel báo cáo tháng: {e}"


# [!] HÀM ĐÃ ĐƯỢC NÂNG CẤP TOÀN DIỆN
def generate_weekly_summary_excel(student_data_list, week_num, teacher_name, class_name, group_name, start_date, end_date, filepath, positive_scorers_list, zero_scorers_list, all_violations_list):
    """
    Tạo file Excel tổng hợp kết quả rèn luyện tuần, có thêm sheet xếp hạng tổ, vinh danh HS xuất sắc VÀ tình hình phát biểu.
    """
    try:
        # 1. Khởi tạo Workbook và các style cơ bản
        wb = openpyxl.Workbook()
        
        font_title = Font(name='Times New Roman', size=16, bold=True)
        font_header = Font(name='Times New Roman', size=12, bold=True)
        font_subheader = Font(name='Times New Roman', size=12, italic=True)
        font_bold = Font(name='Times New Roman', size=11, bold=True)
        font_normal = Font(name='Times New Roman', size=11)
        
        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        
        header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        title_block_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        # === BẮT ĐẦU SHEET 1: TỔNG KẾT CHUNG ===
        ws = wb.active
        ws.title = f"Tong_Ket_Tuan_{week_num}"

        # --- Header ---
        display_group = f" - TỔ {group_name}" if group_name else ""
        title_text = f"BÁO CÁO TỔNG HỢP RÈN LUYỆN TUẦN {week_num}"
        class_text = f"LỚP: {class_name}{display_group}"
        date_text = f"(Từ ngày {start_date.strftime('%d/%m/%Y')} đến {end_date.strftime('%d/%m/%Y')})"
        teacher_text = f"GVCN: {teacher_name}"
        
        ws.merge_cells('A1:E1'); ws['A1'].value = title_text; ws['A1'].font = font_title; ws['A1'].alignment = center_alignment; ws.row_dimensions[1].height = 45
        ws.merge_cells('A2:E2'); ws['A2'].value = class_text; ws['A2'].font = font_header; ws['A2'].alignment = center_alignment
        ws.merge_cells('A3:E3'); ws['A3'].value = date_text; ws['A3'].font = font_subheader; ws['A3'].alignment = center_alignment
        ws.merge_cells('A4:E4'); ws['A4'].value = teacher_text; ws['A4'].font = font_bold; ws['A4'].alignment = center_alignment
        for row in ws['A1:E4']:
            for cell in row: cell.fill = title_block_fill

        # --- Bảng dữ liệu ---
        ws.append([])
        headers = ['STT', 'Họ và Tên', 'Tổ', 'Tổng Điểm', 'Xếp Loại']
        ws.append(headers)
        header_row_num = ws.max_row
        for cell in ws[header_row_num]: cell.font = font_bold; cell.alignment = center_alignment; cell.fill = header_fill; cell.border = thin_border
        for i, student_data in enumerate(student_data_list, 1):
            row_data = [i, student_data['ten'], student_data.get('to') or '', student_data['diem'], student_data['xeploai']]; ws.append(row_data)
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=ws.max_row, column=col_idx); cell.font = font_normal; cell.border = thin_border
                if col_idx in [1, 3, 4, 5]: cell.alignment = center_alignment
                else: cell.alignment = left_alignment
        
        column_widths = {'A': 5, 'B': 35, 'C': 10, 'D': 15, 'E': 15}
        for col, width in column_widths.items(): ws.column_dimensions[col].width = width
        
        ws.print_options.horizontalCentered = True; ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        
        # === BẮT ĐẦU SHEET 2: XẾP HẠNG TỔ ===
        ws_rank = wb.create_sheet("Xep_Hang_To")
        teams = {}; team_ranking = []
        if student_data_list:
            for student in student_data_list:
                team_name = student.get('to')
                if team_name:
                    if team_name not in teams: teams[team_name] = {'total_score': 0, 'count': 0}
                    teams[team_name]['total_score'] += student['diem']; teams[team_name]['count'] += 1
        for name, data in teams.items():
            if data['count'] > 0: avg_score = round(data['total_score'] / data['count'], 2); team_ranking.append({'name': name, 'avg_score': avg_score})
        team_ranking.sort(key=lambda x: x['avg_score'], reverse=True)
        ws_rank.merge_cells('A1:C1'); ws_rank['A1'].value = f"BẢNG XẾP HẠNG THI ĐUA TỔ - TUẦN {week_num}"; ws_rank['A1'].font = font_title; ws_rank['A1'].alignment = center_alignment; ws_rank['A1'].fill = title_block_fill; ws_rank.row_dimensions[1].height = 45
        ws_rank.append([]); rank_headers = ['Xếp Hạng', 'Tên Tổ', 'Điểm Trung Bình']; ws_rank.append(rank_headers)
        for cell in ws_rank[3]: cell.font = font_bold; cell.alignment = center_alignment; cell.fill = header_fill; cell.border = thin_border
        for idx, team in enumerate(team_ranking, 1):
            ws_rank.append([idx, f"Tổ {team['name']}", team['avg_score']])
            for col_idx in range(1, 4):
                cell = ws_rank.cell(row=ws_rank.max_row, column=col_idx); cell.font = font_normal; cell.alignment = center_alignment; cell.border = thin_border
        ws_rank.column_dimensions['A'].width = 15; ws_rank.column_dimensions['B'].width = 25; ws_rank.column_dimensions['C'].width = 25
        ws_rank.print_options.horizontalCentered = True

        # === BẮT ĐẦU SHEET 3: MỜI KHEN THƯỞNG ===
        ws_invite = wb.create_sheet("Moi_Khen_Thuong")
        excellent_students = [s for s in student_data_list if s.get('xeploai') == 'Xuất sắc']
        excellent_students.sort(key=lambda x: x['diem'], reverse=True)
        ws_invite.merge_cells('A1:E1'); ws_invite['A1'].value = f"DANH SÁCH HỌC SINH RÈN LUYỆN XUẤT SẮC - TUẦN {week_num}"; ws_invite['A1'].font = font_title; ws_invite['A1'].alignment = center_alignment; ws_invite['A1'].fill = title_block_fill; ws_invite.row_dimensions[1].height = 45
        ws_invite.append([]); invite_headers = ['STT', 'Họ và Tên', 'Lớp', 'Tổ', 'Điểm Tuần']; ws_invite.append(invite_headers)
        for cell in ws_invite[3]: cell.font = font_bold; cell.alignment = center_alignment; cell.fill = header_fill; cell.border = thin_border
        if not excellent_students: ws_invite.append(['', 'Không có học sinh đạt loại Xuất sắc trong tuần này.'])
        else:
            for idx, student in enumerate(excellent_students, 1):
                ws_invite.append([idx, student['ten'], student['lop'], student.get('to', ''), student['diem']])
                for col_idx in range(1, 6):
                    cell = ws_invite.cell(row=ws_invite.max_row, column=col_idx); cell.font = font_normal; cell.border = thin_border
                    if col_idx != 2: cell.alignment = center_alignment
        ws_invite.column_dimensions['A'].width = 5; ws_invite.column_dimensions['B'].width = 35; ws_invite.column_dimensions['C'].width = 10; ws_invite.column_dimensions['D'].width = 10; ws_invite.column_dimensions['E'].width = 15
        ws_invite.print_options.horizontalCentered = True

        # === BẮT ĐẦU SHEET 4: TỔNG HỢP ĐIỂM CỘNG ===
        ws_speak = wb.create_sheet("Tong_Hop_Diem_Cong")
        ws_speak.merge_cells('A1:C1'); ws_speak['A1'].value = f"BÁO CÁO TỔNG HỢP ĐIỂM CỘNG - TUẦN {week_num}"; ws_speak['A1'].font = font_title; ws_speak['A1'].alignment = center_alignment; ws_speak['A1'].fill = title_block_fill; ws_speak.row_dimensions[1].height = 45
        ws_speak.append([])
        
        ws_speak.merge_cells('A3:C3'); ws_speak['A3'].value = "DANH SÁCH HỌC SINH CÓ ĐIỂM CỘNG"; ws_speak['A3'].font = font_header; ws_speak['A3'].alignment = center_alignment; ws_speak.row_dimensions[3].height = 20
        ws_speak.append([])
        ws_speak.append(['STT', 'Họ và Tên', 'Tổng điểm cộng']);
        start_row = ws_speak.max_row
        for cell in ws_speak[start_row]: cell.font = font_bold; cell.alignment = center_alignment; cell.fill = header_fill; cell.border = thin_border
        for idx, (name, points) in enumerate(positive_scorers_list, 1): # Lặp qua toàn bộ danh sách điểm cộng
            ws_speak.append([idx, name, f"+{points}"])
            for cell in ws_speak[ws_speak.max_row]: cell.font = font_normal; cell.border = thin_border; cell.alignment = center_alignment
            ws_speak.cell(row=ws_speak.max_row, column=2).alignment = left_alignment
        
        last_row_table1 = ws_speak.max_row; row_cursor = last_row_table1 + 3

        ws_speak.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=3)
        cell_title_2 = ws_speak.cell(row=row_cursor, column=1); cell_title_2.value = "DANH SÁCH HỌC SINH CHƯA CÓ ĐIỂM CỘNG"; cell_title_2.font = font_header; cell_title_2.alignment = center_alignment
        ws_speak.row_dimensions[row_cursor].height = 20
        row_cursor += 1

        ws_speak.cell(row=row_cursor, column=1).value = 'STT'; ws_speak.cell(row=row_cursor, column=2).value = 'Họ và Tên'; ws_speak.cell(row=row_cursor, column=3).value = 'Tổng điểm cộng'
        for col_idx in range(1, 4):
            cell = ws_speak.cell(row=row_cursor, column=col_idx); cell.font = font_bold; cell.alignment = center_alignment; cell.fill = header_fill; cell.border = thin_border
        row_cursor += 1

        if not zero_scorers_list: # Lặp qua danh sách học sinh có điểm cộng = 0
            ws_speak.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=3)
            cell = ws_speak.cell(row=row_cursor, column=1); cell.value = 'Tất cả học sinh đều đã có điểm cộng trong tuần. Rất tốt!'; cell.font = font_normal; cell.border = thin_border; cell.alignment = center_alignment
        else:
            for idx, name in enumerate(zero_scorers_list, 1):
                data_row = row_cursor + idx - 1
                ws_speak.cell(row=data_row, column=1).value = idx; ws_speak.cell(row=data_row, column=2).value = name; ws_speak.cell(row=data_row, column=3).value = 0
                for col_idx in range(1, 4):
                    cell = ws_speak.cell(row=data_row, column=col_idx); cell.font = font_normal; cell.border = thin_border
                    if col_idx == 1 or col_idx == 3: cell.alignment = center_alignment
                    else: cell.alignment = left_alignment
                
        ws_speak.column_dimensions['A'].width = 5; ws_speak.column_dimensions['B'].width = 35; ws_speak.column_dimensions['C'].width = 15
        ws_speak.print_options.horizontalCentered = True; ws_speak.page_setup.fitToWidth = 1
        
        # === BẮT ĐẦU SHEET 5: TỔNG HỢP VI PHẠM ===
        ws_violations = wb.create_sheet("TongHop_ViPham")
        ws_violations.merge_cells('A1:F1'); cell_title_violations = ws_violations['A1']; cell_title_violations.value = f"BẢNG TỔNG HỢP VI PHẠM TRONG TUẦN {week_num}"; cell_title_violations.font = font_title; cell_title_violations.alignment = center_alignment; cell_title_violations.fill = title_block_fill; ws_violations.row_dimensions[1].height = 45
        ws_violations.append([])
        violation_headers = ['STT', 'Họ và Tên', 'Lớp', 'Tổ', 'Nội dung Vi phạm', 'Số lần']
        ws_violations.append(violation_headers)
        header_row_num_violations = ws_violations.max_row
        for cell in ws_violations[header_row_num_violations]: cell.font = font_bold; cell.alignment = center_alignment; cell.fill = header_fill; cell.border = thin_border
        if not all_violations_list:
            ws_violations.merge_cells('A4:F4'); cell = ws_violations.cell(row=4, column=1); cell.value = 'Tuyệt vời! Không có học sinh nào vi phạm trong tuần.'; cell.font = font_normal; cell.border = thin_border; cell.alignment = center_alignment
        else:
            for idx, violation in enumerate(all_violations_list, 1):
                row_data = [idx, violation['ten'], violation['lop'], violation.get('to', ''), violation['mo_ta'], violation['count']]
                ws_violations.append(row_data)
                for col_idx in range(1, len(row_data) + 1):
                    cell = ws_violations.cell(row=ws_violations.max_row, column=col_idx); cell.font = font_normal; cell.border = thin_border
                    if col_idx in [1, 3, 4, 6]: cell.alignment = center_alignment
                    else: cell.alignment = left_alignment
        
        ws_violations.column_dimensions['A'].width = 5; ws_violations.column_dimensions['B'].width = 30; ws_violations.column_dimensions['C'].width = 8; ws_violations.column_dimensions['D'].width = 8; ws_violations.column_dimensions['E'].width = 40; ws_violations.column_dimensions['F'].width = 10
        ws_violations.print_options.horizontalCentered = True; ws_violations.page_setup.orientation = ws_violations.ORIENTATION_LANDSCAPE; ws_violations.page_setup.fitToWidth = 1; ws_violations.page_setup.fitToHeight = 0
        
        wb.active = 0
        wb.save(filepath)
        return True, "Xuất file Excel thành công."
        
    except Exception as e:
        return False, f"Lỗi khi tạo file Excel: {e}"

def generate_weekly_logbook_excel(student_data_list, week_num, start_date, end_date, filepath):
    try:
        wb = openpyxl.Workbook()
        from database import lay_tat_ca_danh_muc_su_kien_db
        event_categories = lay_tat_ca_danh_muc_su_kien_db()

        font_bold = Font(name='Calibri', size=11, bold=True)
        font_normal = Font(name='Calibri', size=11)
        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        # === TẠO SHEET 1: SoTheoDoi (để in ra) ===
        ws_print = wb.active
        ws_print.title = "SoTheoDoi"
        ws_print.merge_cells('A1:J1'); ws_print['A1'].value = f"SỔ THEO DÕI RÈN LUYỆN TUẦN {week_num}"; ws_print['A1'].font = Font(name='Calibri', size=16, bold=True); ws_print['A1'].alignment = center_alignment
        ws_print.merge_cells('A2:J2'); ws_print['A2'].value = f"Từ ngày {start_date.strftime('%d/%m/%Y')} đến {end_date.strftime('%d/%m/%Y')}"; ws_print['A2'].font = Font(name='Calibri', size=12, italic=True); ws_print['A2'].alignment = center_alignment
        ws_print.append([])
        headers_print = ['STT', 'Họ và Tên', 'Tổ', 'Thứ 2', 'Thứ 3', 'Thứ 4', 'Thứ 5', 'Thứ 6', 'Thứ 7', 'Ghi chú thêm']
        ws_print.append(headers_print)
        for cell in ws_print[4]: cell.font = font_bold; cell.alignment = center_alignment; cell.border = thin_border
        for i, student in enumerate(student_data_list, 1): ws_print.append([i, student['ten'], student.get('to', '')])
        for row in ws_print.iter_rows(min_row=4, max_row=ws_print.max_row, min_col=1, max_col=10):
            for cell in row:
                cell.border = thin_border
                if cell.column in [1, 3]: cell.alignment = center_alignment
                cell.font = font_normal
        ws_print.column_dimensions['A'].width = 5; ws_print.column_dimensions['B'].width = 30; ws_print.column_dimensions['C'].width = 5
        for col_letter in ['D', 'E', 'F', 'G', 'H', 'I', 'J']: ws_print.column_dimensions[col_letter].width = 15
        for row_idx in range(5, ws_print.max_row + 1): ws_print.row_dimensions[row_idx].height = 25
        ws_print.print_options.horizontalCentered = True; ws_print.page_setup.orientation = ws_print.ORIENTATION_LANDSCAPE; ws_print.page_setup.fitToWidth = 1; ws_print.page_setup.fitToHeight = 0

        # === TẠO SHEET 2: GhiNhan (để nhập liệu và import) ===
        ws_data = wb.create_sheet("GhiNhan")
        headers_data = ['ID Học Sinh (Không sửa)', 'Họ và Tên', 'Ngày Ghi Nhận (dd/mm/yyyy)', 'Mô tả Sự kiện']
        ws_data.append(headers_data)
        for cell in ws_data[1]: cell.font = font_bold
        for student in student_data_list: ws_data.append([student['id'], student['ten']])
        ws_data.column_dimensions['A'].hidden = True; ws_data.column_dimensions['B'].width = 30; ws_data.column_dimensions['C'].width = 25; ws_data.column_dimensions['D'].width = 40

        # === TẠO SHEET 3: DanhMucSuKien (để tham khảo) ===
        ws_events = wb.create_sheet("DanhMucSuKien")
        ws_events.append(['DANH MỤC SỰ KIỆN MẪU (để tham khảo khi nhập liệu)']); ws_events.append(['Tên Sự Kiện'])
        for cat in event_categories: ws_events.append([cat[1]])
        ws_events.column_dimensions['A'].width = 50

        wb.save(filepath)
        return True, "Tải Sổ Ghi nhận Tuần thành công."

    except Exception as e:
        return False, f"Lỗi khi tạo Sổ Ghi nhận Tuần: {e}"